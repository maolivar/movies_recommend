"""Data loading and k-NN functions for the movie recommender (no Streamlit code)."""
import io
import os
import re
import warnings
import urllib.error
import urllib.request

import numpy as np
import pandas as pd
from sklearn.impute import KNNImputer
from sklearn.metrics.pairwise import nan_euclidean_distances

warnings.filterwarnings("ignore", category=RuntimeWarning)   # sklearn matmul on NaN-masked data

DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")

RATING_MAP = {"Me gustó totalmente": 5, "Me gustó": 4, "Ni me gustó, ni disgustó": 3,
              "No me gustó": 2, "No me gustó para nada": 1}
RATING_LOWER = {k.lower(): v for k, v in RATING_MAP.items()}


class SheetError(Exception):
    """The Google Sheet could not be read (bad link, not public, not CSV)."""


def imputer(k):
    return KNNImputer(n_neighbors=k, keep_empty_features=True)


# --------------------------------------------------------------------------- form

def sheet_to_csv_url(link):
    """Turn a Google Sheets link into its CSV export URL."""
    link = link.strip()
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", link)
    if not m:
        raise SheetError("El enlace no parece ser de una hoja de Google Sheets "
                         "(debe contener /spreadsheets/d/...).")
    gid = re.search(r"[#?&]gid=(\d+)", link)
    url = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    return url + (f"&gid={gid.group(1)}" if gid else "")


def read_sheet(link):
    """Download the response sheet as a raw DataFrame."""
    url = sheet_to_csv_url(link)
    try:
        with urllib.request.urlopen(url, timeout=30) as resp:
            content = resp.read()
            ctype = resp.headers.get("Content-Type", "")
    except urllib.error.HTTPError as e:
        if e.code in (401, 403, 404):
            raise SheetError("No se pudo acceder a la hoja (HTTP %d). Compártela como "
                             "«Cualquier persona con el enlace puede ver»." % e.code)
        raise SheetError(f"Error HTTP {e.code} al descargar la hoja.")
    except urllib.error.URLError as e:
        raise SheetError(f"No se pudo conectar con Google: {e.reason}")
    if "text/html" in ctype or content.lstrip()[:1] == b"<":
        raise SheetError("Google devolvió una página web en vez de datos: la hoja no es pública. "
                         "Compártela como «Cualquier persona con el enlace puede ver».")
    return pd.read_csv(io.BytesIO(content))


def movie_title_from_header(header):
    """Text inside the last [...] of a column header (the whole header if there is none)."""
    m = re.search(r"\[([^\[\]]*)\]\s*$", str(header))
    return (m.group(1) if m else str(header)).strip()


def to_ratings(series):
    """Answers -> ratings 1-5. Accepts the Spanish labels of the course form and plain numbers 1-5.

    Returns the ratings (NaN = blank/unrecognized) and the set of non-blank values not recognized.
    """
    s = series.astype("string").str.strip()
    text = s.str.lower().map(RATING_LOWER).astype(float)
    num = pd.to_numeric(s.str.replace(",", ".", regex=False), errors="coerce")
    out = text.fillna(num.where(num.isin([1, 2, 3, 4, 5])))
    bad = s.notna() & s.ne("") & out.isna()
    return out, set(s[bad])


def parse_form(raw, movie_map):
    """Raw sheet -> (DataFrame [timestamp, name, <one column per movie>], movie ids, original titles, notes).

    Layout: first column = timestamp, last = student name, the rest = one column per movie, with the
    title in the last [...] of the header. Movies are shown with the title written in the form; those
    found in movie_map (by form_title) also get their MovieLens id (the returned dict title -> id).
    `originals` maps each displayed name to the title written in the form.
    `notes` lists in Spanish the columns/values that were ignored.
    """
    if raw.shape[1] < 3:
        raise SheetError("La hoja tiene muy pocas columnas: se espera marca temporal, películas y nombre. "
                         "Revisa la estructura requerida del formulario.")
    known = {t.strip().lower(): i for t, i in zip(movie_map["form_title"], movie_map["movieId"])}

    form = pd.DataFrame({"timestamp": pd.to_datetime(raw.iloc[:, 0], errors="coerce"),
                         "name": raw.iloc[:, -1].astype(str).str.strip()})
    movies, ids, originals, notes = {}, {}, {}, []
    for col in raw.columns[1:-1]:
        title = movie_title_from_header(col)
        ratings, bad = to_ratings(raw[col])
        if ratings.notna().sum() == 0:
            why = f" (valores no reconocidos: {', '.join(sorted(bad)[:3])})" if bad else " (sin respuestas)"
            notes.append(f"Columna ignorada: «{title}»{why}")
            continue
        if bad:
            notes.append(f"«{title}»: se ignoraron valores no reconocidos: {', '.join(sorted(bad)[:3])}")
        mid = known.get(title.lower(), np.nan)
        name = title
        base, n = name, 2
        while name in movies:                                  # repeated title: keep both columns
            name, n = f"{base} ({n})", n + 1
        movies[name] = ratings
        ids[name] = mid
        originals[name] = title
    if not movies:
        raise SheetError("No se encontró ninguna columna de películas con calificaciones reconocidas. "
                         "Revisa la estructura requerida del formulario.")
    form = pd.concat([form, pd.DataFrame(movies)], axis=1)
    form = form.dropna(subset=["timestamp"])
    form = form[form["name"].ne("") & form["name"].ne("nan")]
    return form, ids, originals, notes


# --------------------------------------------------------------------------- movie -> MovieLens id mapping

def mapping_template(movie_titles, ids, originals):
    """CSV text (title,movieId) for the professor to fill in; titles as written in the form, known ids prefilled."""
    df = pd.DataFrame({"title": [originals.get(t, t) for t in movie_titles],
                       "movieId": [("" if pd.isna(ids.get(t)) else int(ids[t])) for t in movie_titles]})
    return df.to_csv(index=False)


def read_mapping(content):
    """Uploaded CSV -> DataFrame [title, movieId] (blank ids dropped) and a list of problems."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        df = pd.read_csv(io.StringIO(text), sep=None, engine="python", dtype=str)
    except Exception:
        raise SheetError("No se pudo leer el archivo. Debe ser un CSV con las columnas title y movieId.")
    if df.shape[1] < 2:
        raise SheetError("El archivo debe tener dos columnas: título de la película y movieId.")
    low = {c.strip().lower(): c for c in df.columns}
    tcol = next((low[c] for c in ("title", "form_title", "titulo", "título", "pelicula", "película", "movie", "short_title") if c in low),
                df.columns[0])
    icol = next((low[c] for c in ("movieid", "movie_id", "id") if c in low),
                df.columns[1] if df.columns[1] != tcol else df.columns[0])
    out = pd.DataFrame({"title": df[tcol].astype("string").str.strip(),
                        "raw": df[icol].astype("string").str.strip()})
    out = out.dropna(subset=["title"])
    out["movieId"] = pd.to_numeric(out["raw"], errors="coerce")
    problems = [f"movieId no numérico para «{t}»: {r}" for t, r in zip(out["title"], out["raw"])
                if pd.notna(r) and r != "" and pd.isna(pd.to_numeric(r, errors="coerce"))]
    return out.dropna(subset=["movieId"]).astype({"movieId": int})[["title", "movieId"]], problems


def resolve_ids(movie_titles, builtin_ids, upload, valid_ids, originals=None):
    """MovieLens id per form movie (NaN = unmapped). The upload overrides the built-in map.

    Ids that are not in MovieLens, or used by two movies, are dropped; returns (ids, problems).
    """
    ids = {t: builtin_ids.get(t, np.nan) for t in movie_titles}
    problems = []
    if upload is not None:
        by_title = {t.casefold(): t for t in movie_titles}
        by_title.update({o.casefold(): t for t, o in (originals or {}).items()})   # form title or displayed name
        for t, mid in zip(upload["title"], upload["movieId"]):
            if t.casefold() in by_title:
                ids[by_title[t.casefold()]] = mid
            else:
                problems.append(f"«{t}» no está entre las películas del formulario.")
    seen = {}
    for t in movie_titles:
        mid = ids[t]
        if pd.isna(mid):
            continue
        if mid not in valid_ids:
            problems.append(f"movieId {int(mid)} («{t}») no existe en MovieLens (small): se ignora.")
            ids[t] = np.nan
        elif mid in seen:
            problems.append(f"movieId {int(mid)} repetido en «{seen[mid]}» y «{t}»: se ignora el segundo.")
            ids[t] = np.nan
        else:
            seen[mid] = t
    return pd.Series(ids, dtype=float), problems


def build_ratings(form, start_date, min_ratings):
    """Date filter -> latest answer per name -> drop users with few ratings. Returns users x movies."""
    f = form
    if start_date is not None:
        f = f[f["timestamp"] >= pd.Timestamp(start_date)]
    f = f.sort_values("timestamp", kind="stable").drop_duplicates("name", keep="last")
    ratings = f.set_index("name").drop(columns="timestamp")
    ratings = ratings[ratings.notna().sum(axis=1) >= min_ratings]
    return ratings, f.loc[f["name"].isin(ratings.index), "timestamp"]


# --------------------------------------------------------------------------- k-NN among classmates

def distances(R):
    """nan-Euclidean distance matrix and the matrix of movies rated by both."""
    D = nan_euclidean_distances(R)
    M = (~np.isnan(R)).astype(int)
    return D, M @ M.T


def predict_all(R, k):
    """KNNImputer on the whole users x movies matrix."""
    return imputer(k).fit_transform(R)


def neighbors_behind(R, D, i, c, k):
    """Indices of the k nearest other users (by D) who rated movie c."""
    d = D[i].copy()
    d[i] = np.inf
    order = np.argsort(d, kind="stable")
    return [j for j in order if not np.isnan(R[j, c]) and np.isfinite(d[j])][:k]


def short_name(name):
    parts = name.split()
    return parts[0] if len(parts) <= 1 else f"{parts[0]} {parts[-1][0]}."


# --------------------------------------------------------------------------- MovieLens

def load_movielens():
    """ratings (long), titles (Series by movieId), wide users x movies matrix."""
    ml = pd.read_csv(os.path.join(DATA, "ratings.csv"))
    titles = pd.read_csv(os.path.join(DATA, "movies.csv")).set_index("movieId")["title"]
    wide = ml.pivot_table(index="userId", columns="movieId", values="rating")
    return ml, titles, wide


def candidate_ids(ml, min_ratings=50):
    counts = ml["movieId"].value_counts()
    return counts[counts >= min_ratings].index


def recommend_for_student(wide, form_ids, cand_ids, student_ratings, k, min_common=5):
    """Impute the student's row using MovieLens users who rated >= min_common of the form movies.

    student_ratings: array aligned with form_ids (NaN = not seen). Returns (predictions, neighbors):
    predictions = Series indexed by movieId for the candidate movies that are not form movies;
    neighbors = DataFrame (userId, distance, movies in common) of the pool, closest first.
    Returns (None, None) if the pool is empty.
    """
    form_ids = list(form_ids)
    in_form = set(form_ids)
    cols = form_ids + [m for m in cand_ids if m not in in_form]
    W = wide.reindex(columns=cols)
    pool = W[W[form_ids].notna().sum(axis=1) >= min_common]
    if len(pool) == 0:
        return None, None
    row = np.full((1, len(cols)), np.nan)
    row[0, :len(form_ids)] = student_ratings
    pred = imputer(min(k, len(pool))).fit(pool.to_numpy()).transform(row)[0]
    out = pd.Series(pred[len(form_ids):], index=cols[len(form_ids):])
    d = nan_euclidean_distances(row[:, :len(form_ids)], pool[form_ids].to_numpy())[0]   # same ranking as the imputer
    shared = (pool[form_ids].notna().to_numpy() & ~np.isnan(student_ratings)[None, :]).sum(axis=1)
    near = (pd.DataFrame({"userId": pool.index, "distance": d, "common": shared})
            .sort_values("distance", kind="stable").reset_index(drop=True))
    return out, near


# --------------------------------------------------------------------------- export

def english_titles(movies, movie_map):
    """Display titles for the export: English title if the movie is in movie_map, else as written in the form."""
    eng = {t.strip().lower(): s for t, s in zip(movie_map["form_title"], movie_map["short_title"])}
    return [eng.get(m.strip().lower(), m) for m in movies]


def ratings_to_excel(ratings, titles):
    """Users x movies ratings -> .xlsx bytes (header 'Student' + titles, blue header, blanks = not seen)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Ratings"
    ws.append(["Student"] + list(titles))
    for name, row in zip(ratings.index, ratings.to_numpy(dtype=float)):
        ws.append([name] + [None if np.isnan(v) else int(v) for v in row])

    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", start_color="1F4E79")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[1].height = 45
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = border
            if c.column > 1:
                c.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = max(12, min(40, max(len(str(n)) for n in ratings.index) + 2))
    for j in range(2, len(titles) + 2):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = 12
    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
